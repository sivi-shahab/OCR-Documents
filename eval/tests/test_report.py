import json

import openpyxl

from ocr_eval.numbers import NumberScore
from ocr_eval.report import standard_notes, write_report
from ocr_eval.score import PageRow, summarize


def _row(column, engine, page, recall_matched, **kw):
    number = NumberScore(10, recall_matched, recall_matched, 0, 10 - recall_matched, 0, 0, 0,
                         dropped_examples=["1234567"])  # fmt: skip
    fields = dict(
        document="doc", kind="digital", number=number, dice=0.8, english=None, columns=None, totals=None,
        consensus=None, empty=False, truncated=False, error=None, seconds=1.0,
    )  # fmt: skip
    fields.update(kw)
    return PageRow(column=column, engine=engine, page=page, **fields)


def test_report_files(tmp_path):
    rows = [_row("digital", "unlimited", p, p) for p in range(1, 25)]
    rows += [_row("digital", "docling", 1, 10), _row("synthetic_scan", "docling", 1, 7)]
    rows += [
        _row("real_scan", "mistral", 3, 0, number=None, kind="scan", consensus=NumberScore(4, 3, 3, 0, 1, 0, 0, 0))
    ]
    metas = {
        ("original", "doc", "unlimited"): {
            "model": "Unlimited-OCR",
            "concurrency": 32,
            "wall_seconds": 10,
            "pages": {},
        },
        ("original", "doc", "docling"): {"model": "docling (EasyOCR + TableFormer)", "version": "2.143.0"},
        ("original", "doc", "mistral"): {"model": "mistral-ocr-latest"},
    }
    (tmp_path / "classification.json").write_text(json.dumps({"doc": {"1": "digital", "3": "scan"}}))
    notes = standard_notes(
        metas, {"azure": "OCR_AZURE_DI_API_KEY kosong"}, excluded_reference_amounts=7, gpu_hour_rate=0.95
    )
    write_report(rows, summarize(rows, metas, 0.95), tmp_path, notes)

    book = openpyxl.load_workbook(tmp_path / "ocr_eval.xlsx")
    assert book.sheetnames == ["Summary", "Pages", "Worst pages", "Classification", "Notes"]
    worst = list(book["Worst pages"].iter_rows(values_only=True))
    unlimited_worst = [r for r in worst[1:] if r[1] == "unlimited"]
    assert len(unlimited_worst) == 20
    assert any("1234567" in str(cell) for cell in worst[1])
    summary = list(book["Summary"].iter_rows(values_only=True))
    header = summary[0]
    docling_digital = next(r for r in summary[1:] if r[:2] == ("digital", "docling"))
    for name in ("column_placement", "column_end_to_end"):
        assert docling_digital[header.index(name)] == "n/a (acuan)"
    classification = list(book["Classification"].iter_rows(values_only=True))
    assert ("doc", 3, "scan") in classification

    md = (tmp_path / "summary.md").read_text(encoding="utf-8")
    for note in notes:
        assert note in md
    notes_sheet = [r[0] for r in book["Notes"].iter_rows(values_only=True)]
    for note in notes:
        assert note in notes_sheet
    text = "\n".join(notes)
    for required in ("`digital`", "`synthetic_scan`", "bias", "utilisasi penuh", "biaya operasional", "AGPL",
                     "legal", "azure: OCR_AZURE_DI_API_KEY kosong", "Unlimited-OCR", "docling (EasyOCR + TableFormer) 2.143.0",
                     "mistral-ocr-latest", "7 angka acuan",
                     "`column_placement`", "`column_end_to_end`", "`median_seconds`/`p95_seconds`",
                     "hanya untuk Unlimited-OCR", "`truncated`", "finish_reason=length", "satu-satunya"):  # fmt: skip
        assert required in text, required


def test_illegal_characters_are_stripped(tmp_path):
    rows = [_row("digital", "unlimited", 1, 5, error="engine said \x1b[31mboom\x00")]
    rows[0].number.dropped_examples = ["12\x1b34567"]
    write_report(rows, summarize(rows, {}, 0.95), tmp_path, ["catatan \x1b"])
    book = openpyxl.load_workbook(tmp_path / "ocr_eval.xlsx")
    pages = list(book["Pages"].iter_rows(values_only=True))
    assert "engine said [31mboom" in pages[1]
    worst = list(book["Worst pages"].iter_rows(values_only=True))
    assert "1234567" in worst[1] and "engine said [31mboom" in worst[1]
    assert [r[0] for r in book["Notes"].iter_rows(values_only=True)] == ["note", "catatan "]
